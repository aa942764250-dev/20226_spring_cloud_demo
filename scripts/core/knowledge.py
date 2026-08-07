import atexit
import hashlib
import re
import sqlite3
import uuid
from pathlib import Path

import fitz
from docx import Document
from qdrant_client import QdrantClient, models
from sentence_transformers import SentenceTransformer


ROOT = Path(__file__).parent
DOCUMENTS_DIR = ROOT / "data" / "documents"
SQLITE_PATH = ROOT / "data" / "sqlite" / "knowledge.db"
COLLECTION = "local_knowledge"
MODEL_NAME = "BAAI/bge-small-zh-v1.5"
VECTOR_SIZE = 512
CHUNK_SIZE = 800
CHUNK_OVERLAP = 120
_model = None
_qdrant = None


def _db() -> sqlite3.Connection:
    SQLITE_PATH.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(SQLITE_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA journal_mode=WAL")
    connection.execute(
        """CREATE TABLE IF NOT EXISTS chunks (
        id TEXT PRIMARY KEY, source TEXT NOT NULL, title TEXT NOT NULL,
        page INTEGER, content TEXT NOT NULL, content_hash TEXT NOT NULL)"""
    )
    connection.execute(
        "CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(id UNINDEXED, content)"
    )
    return connection


def _client() -> QdrantClient:
    global _qdrant
    if _qdrant is None:
        _qdrant = QdrantClient(path=str(ROOT / "data" / "qdrant"))
    return _qdrant


@atexit.register
def _close_client() -> None:
    if _qdrant is not None:
        _qdrant.close()


def _ensure_collection(client: QdrantClient) -> None:
    if client.collection_exists(COLLECTION):
        return
    client.create_collection(
        collection_name=COLLECTION,
        vectors_config=models.VectorParams(
            size=VECTOR_SIZE,
            distance=models.Distance.COSINE,
            on_disk=True,
        ),
        hnsw_config=models.HnswConfigDiff(on_disk=True),
        on_disk_payload=True,
    )


def _embedding_model() -> SentenceTransformer:
    global _model
    if _model is None:
        _model = SentenceTransformer(MODEL_NAME, device="cpu")
    return _model


def unload_embedding_model() -> None:
    global _model
    _model = None


def _read_document(path: Path) -> list[tuple[int | None, str]]:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return [(None, path.read_text(encoding="utf-8", errors="ignore"))]
    if suffix == ".pdf":
        pdf = fitz.open(path)
        try:
            return [(index + 1, page.get_text("text")) for index, page in enumerate(pdf)]
        finally:
            pdf.close()
    if suffix == ".docx":
        return [(None, "\n".join(p.text for p in Document(path).paragraphs))]
    raise ValueError(f"unsupported file type: {path.suffix}")


def _chunks(text: str) -> list[str]:
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    result, start = [], 0
    while start < len(text):
        end = min(len(text), start + CHUNK_SIZE)
        if end < len(text):
            boundary = max(text.rfind("\n", start, end), text.rfind("。", start, end))
            if boundary > start + CHUNK_SIZE // 2:
                end = boundary + 1
        result.append(text[start:end].strip())
        if end == len(text):
            break
        start = max(end - CHUNK_OVERLAP, start + 1)
    return [item for item in result if item]


def index_document(file_path: str) -> dict:
    path = Path(file_path).resolve()
    if not path.is_relative_to(DOCUMENTS_DIR.resolve()):
        raise PermissionError(f"file must be inside {DOCUMENTS_DIR}")
    if not path.is_file():
        raise FileNotFoundError(path)

    source = str(path.relative_to(DOCUMENTS_DIR)).replace("\\", "/")
    entries = []
    for page, text in _read_document(path):
        for content in _chunks(text):
            position = len(entries)
            chunk_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"{source}:{position}:{content}"))
            entries.append((chunk_id, source, path.stem, page, content))
    if not entries:
        raise ValueError("no readable text found; scanned PDFs require OCR")

    client = _client()
    _ensure_collection(client)
    client.delete(
        COLLECTION,
        models.FilterSelector(filter=models.Filter(must=[models.FieldCondition(
            key="source", match=models.MatchValue(value=source)
        )])),
    )
    with _db() as db:
        old_ids = [row[0] for row in db.execute("SELECT id FROM chunks WHERE source = ?", (source,))]
        if old_ids:
            db.executemany("DELETE FROM chunks_fts WHERE id = ?", [(item,) for item in old_ids])
        db.execute("DELETE FROM chunks WHERE source = ?", (source,))

        vectors = _embedding_model().encode(
            [entry[4] for entry in entries], normalize_embeddings=True, show_progress_bar=False
        ).tolist()
        points = []
        for entry, vector in zip(entries, vectors):
            chunk_id, item_source, title, page, content = entry
            digest = hashlib.sha256(content.encode("utf-8")).hexdigest()
            db.execute(
                "INSERT INTO chunks VALUES (?, ?, ?, ?, ?, ?)",
                (chunk_id, item_source, title, page, content, digest),
            )
            db.execute("INSERT INTO chunks_fts VALUES (?, ?)", (chunk_id, content))
            points.append(models.PointStruct(
                id=chunk_id, vector=vector,
                payload={"source": item_source, "title": title, "page": page},
            ))
        client.upsert(COLLECTION, points=points, wait=True)
    return {"source": source, "chunks": len(entries)}


def _keyword_hits(query: str, limit: int) -> list[str]:
    terms = re.findall(r"[\w\u4e00-\u9fff]+", query)
    if not terms:
        return []
    expression = " OR ".join(f'"{term}"' for term in terms)
    with _db() as db:
        return [row[0] for row in db.execute(
            "SELECT id FROM chunks_fts WHERE chunks_fts MATCH ? LIMIT ?", (expression, limit)
        )]


def search(query: str, top_k: int = 5) -> list[dict]:
    query = query.strip()
    if not query:
        raise ValueError("query must not be empty")
    top_k = max(1, min(top_k, 10))
    client = _client()
    _ensure_collection(client)
    vector = _embedding_model().encode(query, normalize_embeddings=True).tolist()
    vector_hits = client.query_points(COLLECTION, query=vector, limit=top_k * 3).points
    keyword_ids = _keyword_hits(query, top_k * 3)
    scores = {}
    for rank, hit in enumerate(vector_hits, 1):
        scores[str(hit.id)] = scores.get(str(hit.id), 0) + 1 / (60 + rank)
    for rank, chunk_id in enumerate(keyword_ids, 1):
        scores[chunk_id] = scores.get(chunk_id, 0) + 1 / (60 + rank)
    if not scores:
        return []
    ids = sorted(scores, key=scores.get, reverse=True)[:top_k]
    placeholders = ",".join("?" for _ in ids)
    with _db() as db:
        rows = {row["id"]: row for row in db.execute(
            f"SELECT id, source, title, page, content FROM chunks WHERE id IN ({placeholders})", ids
        )}
    return [{
        "chunk_id": item_id, "source": rows[item_id]["source"],
        "title": rows[item_id]["title"], "page": rows[item_id]["page"],
        "score": round(scores[item_id], 5), "content": rows[item_id]["content"][:1500],
    } for item_id in ids if item_id in rows]


def get_chunk(chunk_id: str) -> dict:
    with _db() as db:
        row = db.execute(
            "SELECT id, source, title, page, content FROM chunks WHERE id = ?", (chunk_id,)
        ).fetchone()
    if row is None:
        raise KeyError("chunk not found")
    return dict(row)


def list_sources() -> list[str]:
    with _db() as db:
        return [row[0] for row in db.execute("SELECT DISTINCT source FROM chunks ORDER BY source")]
